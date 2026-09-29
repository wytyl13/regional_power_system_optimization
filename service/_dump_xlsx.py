# -*- coding: utf-8 -*-
"""dump 旧系统输入/输出 Excel 的结构：sheet 名、维度、表头、前几行"""
import sys
import openpyxl

def dump(path, max_rows=6, max_cols=40):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"=== {path.split(chr(92))[-1]} ===")
    print(f"sheets ({len(wb.sheetnames)}): {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n--- [{name}] dims={ws.max_row}x{ws.max_column} ---")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True)):
            vals = []
            for v in row:
                if v is None:
                    vals.append("")
                else:
                    s = str(v)
                    if len(s) > 14:
                        s = s[:13] + "…"
                    vals.append(s)
            # 去掉行尾空
            while vals and vals[-1] == "":
                vals.pop()
            print(f"r{i+1}: {'|'.join(vals)}")
        n = 0
        for _ in ws.iter_rows(values_only=True):
            n += 1
        if n > max_rows:
            print(f"... ({n - max_rows} more rows)")
    wb.close()

if __name__ == "__main__":
    for p in sys.argv[1:]:
        dump(p)
